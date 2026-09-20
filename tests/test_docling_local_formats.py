"""Local format conversion follows Docling's declared pipelines without PDF coercion."""

import pytest

from openreading.adapters.docling_local import DoclingLocalAdapter
from openreading.adapters.docling_local.client import LocalDoclingClient
from openreading.adapters.docling_local.config import LocalDoclingConfig
from openreading.adapters.docling_local.projection import project_document
from openreading.types.request import Outputs


@pytest.mark.parametrize(
    "suffix,body",
    [
        ("md", "# FORMAT CHECK\n\nLiteral café paragraph."),
        ("html", "<html><body><h1>FORMAT CHECK</h1><p>Literal café paragraph.</p></body></html>"),
        ("csv", "name,value\nFORMAT CHECK,Literal café paragraph.\n"),
    ],
)
def test_model_free_formats_preserve_provider_text(tmp_path, monkeypatch, suffix, body):
    pytest.importorskip("docling")
    from openreading.adapters.docling_local.pipeline import create_converter

    monkeypatch.setattr(LocalDoclingConfig, "validate_assets", lambda self: {})
    path = tmp_path / f"sample.{suffix}"
    path.write_text(body)
    client = LocalDoclingClient(LocalDoclingConfig(tmp_path))
    raw = client.convert_path(path)
    expected = create_converter(client.config).convert(path).document.export_to_text()
    value, origins = project_document(raw, Outputs())
    assert "FORMAT CHECK" in value.document.text
    assert "Literal café paragraph." in value.document.text
    assert value.document.text == expected
    assert value.document.page_count is None
    assert origins == {}
    assert any(w.code == "page_attribution_unavailable" for w in value.warnings)
    assert suffix in DoclingLocalAdapter().descriptor.capabilities.input_formats


def test_unified_api_preserves_the_input_format(tmp_path, monkeypatch):
    from openreading import run
    from openreading.adapters.docling_local import client

    monkeypatch.setattr(LocalDoclingConfig, "validate_assets", lambda self: {})
    monkeypatch.setenv("DOCLING_LOCAL_ASSETS", str(tmp_path))
    monkeypatch.setattr(client, "_shared", client._SharedClient())
    path = tmp_path / "sample.md"
    path.write_text("# FORMAT CHECK\n\nLiteral paragraph.")
    for source, options in [(str(path), {}), (path.read_bytes(), {"mime_type": "text/markdown"})]:
        result = run(source, backend="docling_local", **options)
        assert "FORMAT CHECK" in result["document"]["text"]
        assert "page_count" not in result["document"]
    client._shared._discard()


@pytest.mark.parametrize("suffix", ["docx", "pptx", "xlsx"])
def test_office_content_does_not_claim_physical_pdf_pages(tmp_path, monkeypatch, suffix):
    monkeypatch.setattr(LocalDoclingConfig, "validate_assets", lambda self: {})
    path = tmp_path / f"sample.{suffix}"
    if suffix == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph("FORMAT CHECK")
        document.save(path)
    elif suffix == "pptx":
        from pptx import Presentation

        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[0])
        slide.shapes.title.text = "FORMAT CHECK"
        document.save(path)
    else:
        from openpyxl import Workbook

        document = Workbook()
        document.active["A1"] = "FORMAT CHECK"
        document.active["B1"] = "42"
        document.save(path)
    raw = LocalDoclingClient(LocalDoclingConfig(tmp_path)).convert_path(path)
    value, origins = project_document(raw, Outputs())
    assert "FORMAT CHECK" in value.document.text
    assert value.document.page_count is None
    assert origins == {}
    assert any(w.code == "page_attribution_unavailable" for w in value.warnings)
    assert any("FORMAT CHECK" in (b.text or "") for b in value.document.pages[0].blocks)


def test_unpaginated_table_keeps_reported_cells_and_exact_block_reference():
    from openreading.artifacts.passages import iter_passages

    raw = {
        "unpaginated": True,
        "document_text": "Name  Amount\nACME  42",
        "pages": {"1": {"size": {"width": 2, "height": 2}}},
        "items": [
            {
                "self_ref": "#/tables/0",
                "label": "table",
                "data": {
                    "table_cells": [
                        {"text": "ACME", "start_row_offset_idx": 0, "start_col_offset_idx": 0},
                        {"text": "42", "start_row_offset_idx": 0, "start_col_offset_idx": 1},
                    ]
                },
            }
        ],
    }
    value, origins = project_document(raw, Outputs(tables="cells"))
    block = value.document.pages[0].blocks[0]
    assert block.id == "#/tables/0"
    assert block.table.rows == [["ACME", "42"]]
    assert block.bbox is None and origins == {}
    passage = next(iter_passages(value))
    assert passage.page is None
    assert passage.source_block_id == "#/tables/0"
    assert passage.source_pointer == "/document/pages/0/blocks/0/text"


def test_selection_extensions_follow_the_configured_provider_formats():
    from docling.datamodel.base_models import FormatToExtensions, InputFormat

    from openreading.adapters.docling_local import formats

    expected = {
        extension.lower()
        for name in formats.INPUT_FORMATS
        for extension in FormatToExtensions[InputFormat(name)]
    }
    assert set(formats.selection_extensions()) == expected
    assert {"docx", "png", "csv", "txt", "pdf"} <= expected
    assert "mp4" not in expected


def test_image_coordinates_keep_the_provider_point_units():
    from docling.backend.image_backend import _ImagePageBackend
    from PIL import Image

    # The pinned backend rescales pixel dimensions by DPI before emitting provenance.
    backend = _ImagePageBackend(Image.new("RGB", (400, 200)), 0, (144, 144))
    assert backend.get_size().width == 200
    assert backend.get_size().height == 100
    raw = {
        "page_unit": "pixel",
        "pages": {"1": {"size": {"width": 200, "height": 100}}},
        "items": [
            {
                "label": "text",
                "text": "image text",
                "prov": [
                    {
                        "page_no": 1,
                        "charspan": [0, 10],
                        "bbox": {"l": 20, "t": 10, "r": 120, "b": 30, "coord_origin": "TOPLEFT"},
                    }
                ],
            }
        ],
        "page_origins": {"1": "ocr"},
    }
    value, origins = project_document(raw, Outputs())
    page = value.document.pages[0]
    assert page.unit.value == "pdf_point"
    assert page.blocks[0].bbox.bbox_native.unit.value == "pdf_point"
    assert page.blocks[0].bbox.x == 0.1
    assert origins == {1: "ocr"}


def test_retained_worker_preserves_model_free_table_cells(tmp_path, monkeypatch):
    import json

    from openreading.artifacts.worker import extract

    monkeypatch.setattr(LocalDoclingConfig, "validate_assets", lambda self: {})
    (tmp_path / "source.csv").write_text("name,value\nACME,42\n")
    extract(
        {
            "directory": str(tmp_path),
            "source_file": "source.csv",
            "pages": None,
            "available": None,
            "extraction_bytes": None,
            "docling": LocalDoclingConfig(tmp_path).wire(),
            "expected_assets": {},
        }
    )
    content = json.loads((tmp_path / "response.json").read_bytes())
    table = next(
        block for block in content["document"]["pages"][0]["blocks"] if block["type"] == "table"
    )
    assert table["table"]["rows"] == [["name", "value"], ["ACME", "42"]]
    assert content["channel_provenance"]["table_cells"] == "native"


@pytest.mark.parametrize(
    "suffix,body",
    [
        (
            "md",
            "# FORMAT CHECK\n\n![remote](https://example.invalid/image.png)\n![local](file:///secret.png)",
        ),
        (
            "html",
            '<html><body><h1>FORMAT CHECK</h1><img src="https://example.invalid/image.png"><img src="file:///secret.png"></body></html>',
        ),
    ],
)
def test_linked_resources_are_not_fetched(tmp_path, monkeypatch, suffix, body):
    import socket

    from docling.backend.abstract_backend import AbstractDocumentBackend

    monkeypatch.setattr(LocalDoclingConfig, "validate_assets", lambda self: {})
    original = AbstractDocumentBackend.__init__
    observed = []

    def checked(self, *args, **kwargs):
        original(self, *args, **kwargs)
        observed.append(self.options)
        assert not self.options.enable_remote_fetch
        assert not self.options.enable_local_fetch
        assert not getattr(self.options, "fetch_images", False)

    def refused(*args, **kwargs):
        raise AssertionError("Conversion attempted network access")

    monkeypatch.setattr(AbstractDocumentBackend, "__init__", checked)
    monkeypatch.setattr(socket.socket, "connect", refused)
    source = tmp_path / f"source.{suffix}"
    source.write_text(body)
    raw = LocalDoclingClient(LocalDoclingConfig(tmp_path)).convert_path(source)
    assert "FORMAT CHECK" in raw["document_text"]
    assert observed


def test_model_free_projection_preserves_flags_and_requested_channels():
    import copy

    raw = {
        "pages": {},
        "document_text": "12. Literal item",
        "partial": True,
        "omitted_furniture_items": 1,
        "items": [
            {
                "self_ref": "#/texts/0",
                "label": "list_item",
                "orig": "",
                "text": "12. Literal item",
                "children": [{"cref": "#/texts/1"}],
            },
            {"self_ref": "#/pictures/0", "label": "picture"},
        ],
    }
    saved = copy.deepcopy(raw)
    value, _ = project_document(raw, Outputs(markdown=True, typed_fields=True))
    assert value.status.state.value == "partial"
    block = value.document.pages[0].blocks[0]
    assert block.text == "12. Literal item" and block.children == ["#/texts/1"]
    assert value.document.pages[0].blocks[1].type.value == "figure"
    assert {w.code for w in value.warnings} == {
        "page_attribution_unavailable",
        "partial_conversion",
        "furniture_text_omitted",
        "unsupported_channel",
    }
    assert {w.field for w in value.warnings if w.code == "unsupported_channel"} == {
        "markdown",
        "typed_fields",
    }
    value, _ = project_document(raw, Outputs(text=False, blocks=False))
    assert value.document.text is None and value.document.pages is None
    assert raw == saved
